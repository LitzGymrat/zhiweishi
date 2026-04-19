from __future__ import annotations

import csv
from dataclasses import dataclass
from pathlib import Path

import docx2txt
from openpyxl import load_workbook
from pypdf import PdfReader

from src.corpus_manager import infer_document_context


@dataclass(slots=True)
class RawDocument:
    path: Path
    source_label: str
    content: str
    doc_category: str
    device_name: str


def read_text_file(file_path: Path) -> str:
    return file_path.read_text(encoding="utf-8", errors="ignore")


def read_pdf_file(file_path: Path) -> str:
    reader = PdfReader(str(file_path))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def read_docx_file(file_path: Path) -> str:
    return docx2txt.process(str(file_path)) or ""


def read_table_file(file_path: Path) -> str:
    if file_path.suffix.lower() == ".csv":
        with file_path.open("r", encoding="utf-8", errors="ignore", newline="") as handle:
            reader = csv.reader(handle)
            return "\n".join(",".join(cell for cell in row) for row in reader)

    workbook = load_workbook(filename=file_path, read_only=True, data_only=True)
    lines: list[str] = []
    for worksheet in workbook.worksheets:
        lines.append(f"[Sheet] {worksheet.title}")
        for row in worksheet.iter_rows(values_only=True):
            cleaned = ["" if item is None else str(item) for item in row]
            if any(cell.strip() for cell in cleaned):
                lines.append(",".join(cleaned))
    return "\n".join(lines)


def load_document(file_path: Path) -> RawDocument:
    suffix = file_path.suffix.lower()
    if suffix in {".txt", ".md"}:
        content = read_text_file(file_path)
    elif suffix == ".pdf":
        content = read_pdf_file(file_path)
    elif suffix == ".docx":
        content = read_docx_file(file_path)
    elif suffix in {".csv", ".xlsx"}:
        content = read_table_file(file_path)
    else:
        raise ValueError(f"不支持的文件类型：{file_path.suffix}")

    document_context = infer_document_context(file_path)
    return RawDocument(
        path=file_path,
        source_label=file_path.name,
        content=content.strip(),
        doc_category=document_context["category_label"],
        device_name=document_context["device_name"],
    )
